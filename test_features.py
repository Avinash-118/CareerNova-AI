"""Integration coverage for the extended workspace. Uses disposable accounts only."""
import io
import json
import os
import secrets
import unittest
from datetime import datetime,timedelta,timezone
from unittest.mock import patch
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import text
import backend
import auth_routes
import oauth_routes
import alert_worker
import platform_routes

class ExtendedWorkspace(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.context=TestClient(backend.app);cls.context.__enter__()
        cls.accounts=[];cls.run_ids=[];cls.domains=[];cls.report_ids=[]

    @classmethod
    def tearDownClass(cls):
        with backend.engine.begin() as conn:
            for uid in cls.accounts:
                conn.execute(text('DELETE FROM public.career_job_tracker WHERE user_id=:u'),{'u':str(uid)})
                conn.execute(text('DELETE FROM public.users WHERE id=:u'),{'u':uid})
            for rid in cls.run_ids:conn.execute(text('DELETE FROM public.worker_runs WHERE id=:id'),{'id':rid})
            for domain in cls.domains:conn.execute(text("DELETE FROM public.employer_reviews WHERE domain=:d"),{"d":domain})
            for report_id in cls.report_ids:conn.execute(text("DELETE FROM public.job_reports WHERE id=:id"),{"id":report_id})
        cls.context.__exit__(None,None,None)

    def account(self):
        client=TestClient(backend.app,headers={'X-CareerNova-Client':'web'},follow_redirects=False)
        email='cn-feature-'+secrets.token_hex(8)+'@example.com'
        result=client.post('/api/auth/register',json={'full_name':'Workspace Tester','email':email,'password':'Safe-test-password-2026!'})
        self.assertEqual(result.status_code,201,result.text)
        user=client.get('/api/auth/me').json();self.accounts.append(user['id'])
        return client,user

    def test_workspace_isolation_and_lifecycle(self):
        a,user=self.account();b,other=self.account()
        doc=Document();doc.add_heading('Experience');doc.add_paragraph('Built Python SQL tools for 100 users.');doc.add_heading('Education');doc.add_paragraph('University degree')
        file=io.BytesIO();doc.save(file)
        upload=a.post('/api/resumes',data={'name':'Engineering'},files={'file':('resume.docx',file.getvalue())})
        self.assertEqual(upload.status_code,201,upload.text);rid=upload.json()['id']
        self.assertEqual(b.get(f'/api/resumes/{rid}/download').status_code,404)
        self.assertEqual(b.get('/api/resumes').json()['resumes'],[])
        self.assertEqual(a.get(f'/api/resumes/{rid}/download').content,file.getvalue())
        description='Build Python and SQL services with Docker and React for customers.'
        coach=a.post('/api/resume-coach',json={'resume_id':rid,'job_description':description})
        self.assertEqual(coach.status_code,200,coach.text);self.assertIn('Docker',coach.json()['missing_skills'])
        self.assertIn('[verified result]',coach.json()['bullet_suggestions'][0]['suggested'])
        self.assertEqual(b.post('/api/resume-coach',json={'resume_id':rid,'job_description':description}).status_code,404)
        with patch('faiss_search.semantic_search',return_value=[{'id':'x','title':'Engineer','description':description,'faiss_score':0.8}]):
            self.assertEqual(a.post(f'/api/resumes/{rid}/matches').status_code,200)
        class FakeModel:
            def encode(self,*args,**kwargs):
                import numpy as np
                return np.array([[1.0,0.0],[0.8,0.6]])
        with patch('faiss_search.load_search_engine',return_value=(None,None,FakeModel())):
            explained=a.post('/api/match-explanation',json={'resume_id':rid,'job_description':description})
            self.assertEqual(explained.status_code,200,explained.text);self.assertIn('70%',explained.json()['formula'])
        practice=a.post('/api/interviews',json={'job_title':'Python Engineer','job_description':description})
        self.assertEqual(practice.status_code,201,practice.text);pid=practice.json()['id']
        self.assertTrue(any(q['category']=='Python' for q in practice.json()['questions']))
        self.assertEqual(a.patch(f'/api/interviews/{pid}',json={'answers':{'intro':'My actual experience'}}).status_code,200)
        self.assertEqual(b.patch(f'/api/interviews/{pid}',json={'answers':{}}).status_code,404)
        self.assertEqual(a.patch(f'/api/interviews/{pid}',json={'answers':{'invented':'invalid'}}).status_code,400)
        a.post('/api/dashboard/save',json={'job_id':'ext-test','job_title':'Engineer'})
        tid=a.get('/api/dashboard/jobs').json()['jobs'][0]['id']
        details={'resume_id':rid,'salary':'$90,000–110,000','description':description}
        self.assertEqual(a.patch(f'/api/applications/{tid}/details',json=details).status_code,200)
        self.assertEqual(b.patch(f'/api/applications/{tid}/details',json=details).status_code,404)
        self.assertEqual(a.get('/api/dashboard/jobs').json()['jobs'][0]['resume_id'],rid)
        reminder=a.post('/api/reminders',json={'title':'Interview tomorrow','kind':'Interview','tracker_id':tid,'due_at':(datetime.now(timezone.utc)+timedelta(hours=2)).isoformat()})
        self.assertEqual(reminder.status_code,201,reminder.text);remid=reminder.json()['id']
        self.assertEqual(b.patch(f'/api/reminders/{remid}',json={'completed':True}).status_code,404)
        self.assertEqual(a.patch(f'/api/reminders/{remid}',json={'completed':True}).status_code,200)
        report=a.post('/api/reports',json={'job_id':'ext-test','title':'Engineer','url':'https://example.com/job','reason':'Expired','details':'The employer says this opening is closed.'})
        self.assertEqual(report.status_code,201,report.text)
        self.report_ids.append(report.json()['id'])
        self.assertEqual(b.get('/api/reports').json()['reports'],[])
        self.assertEqual(a.get('/api/admin/overview').status_code,403)
        with backend.engine.begin() as conn:conn.execute(text('UPDATE public.users SET email_verified=TRUE WHERE id=:u'),{'u':user['id']})
        domain='cn-test-'+secrets.token_hex(4)+'.example.com'
        self.domains.append(domain)
        with patch.dict(os.environ,{'ADMIN_EMAILS':user['email']}):
            self.assertEqual(a.get('/api/admin/overview').status_code,200)
            self.assertEqual(a.patch(f"/api/admin/reports/{report.json()['id']}",json={'status':'Resolved'}).status_code,200)
            self.assertEqual(a.post('/api/admin/employers',json={'domain':domain,'company':'Test Employer','status':'Verified','evidence':'Test-only evidence for integration checks.'}).status_code,200)
        exported=a.get('/api/privacy/export');self.assertEqual(exported.status_code,200)
        export=exported.json();self.assertEqual(export['account']['id'],user['id']);self.assertNotIn('password_hash',exported.text);self.assertEqual(len(export['resume_versions']),1)
        sessions=a.get('/api/privacy/sessions').json()['sessions'];self.assertTrue(sessions[0]['current'])
        self.assertEqual(b.delete('/api/privacy/sessions/'+sessions[0]['id']).status_code,200)
        self.assertEqual(a.get('/api/auth/me').status_code,200)
        self.assertEqual(a.delete(f'/api/resumes/{rid}').status_code,200)
        self.assertIsNone(a.get('/api/dashboard/jobs').json()['jobs'][0]['resume_id'])
        self.assertEqual(a.post('/api/privacy/delete-account',json={'confirmation':'DELETE','password':'wrong'}).status_code,403)
        self.assertEqual(a.post('/api/privacy/delete-account',json={'confirmation':'DELETE','password':'Safe-test-password-2026!'}).status_code,200)
        self.assertEqual(a.get('/api/auth/me').status_code,401)
        with backend.engine.connect() as conn:
            self.assertEqual(conn.execute(text('SELECT COUNT(*) FROM public.resume_versions WHERE user_id=:u'),{'u':user['id']}).scalar_one(),0)
            self.assertIsNone(conn.execute(text('SELECT user_id FROM public.job_reports WHERE id=:id'),{'id':report.json()['id']}).scalar_one())

    def test_worker_verification_digest_retry_and_optout(self):
        a,user=self.account()
        a.post('/api/alerts',json={'query':'Python'})
        a.patch('/api/notification-preferences',json={'email_digest':True})
        self.assertFalse(a.get('/api/notification-preferences').json()['email_verified'])
        a.post('/api/verify-email/request')
        with backend.engine.connect() as conn:
            body=conn.execute(text('SELECT body FROM public.mail_outbox WHERE user_id=:u'),{'u':user['id']}).scalar_one()
        token=body.split('token=')[1]
        self.assertEqual(a.post('/api/verify-email',json={'token':token}).status_code,200)
        self.assertEqual(a.post('/api/verify-email',json={'token':token}).status_code,400)
        job={'id':'worker-test','title':'Python Engineer','company':'Test','location':'Remote','description':'Python','employment_type':'full-time','url':'https://example.com/job'}
        with patch.object(platform_routes,'fetch_live_jobs',return_value=[job]),patch.dict(os.environ,{'SMTP_EMAIL':'test@example.com','SMTP_APP_PASSWORD':'test-only'}),patch.object(alert_worker,'send_email',return_value=True) as send:
            result=alert_worker.run_cycle(user_ids=[user['id']]);self.run_ids.append(result['run_id'])
            self.assertEqual(result['status'],'Completed');self.assertEqual(result['matches'],1);self.assertEqual(send.call_count,2)
            second=alert_worker.run_cycle(user_ids=[user['id']]);self.run_ids.append(second['run_id']);self.assertEqual(second['delivered'],0)
        # A new verification message uses the retry path on SMTP failure.
        a.post('/api/verify-email/request')
        with patch.object(platform_routes,'fetch_live_jobs',return_value=[job]),patch.dict(os.environ,{'SMTP_EMAIL':'test@example.com','SMTP_APP_PASSWORD':'test-only'}),patch.object(alert_worker,'send_email',side_effect=RuntimeError('test failure')):
            result=alert_worker.run_cycle(user_ids=[user['id']]);self.run_ids.append(result['run_id'])
            self.assertEqual(result['status'],'Completed')
        with backend.engine.connect() as conn:
            row=conn.execute(text("SELECT status,attempts FROM public.mail_outbox WHERE user_id=:u ORDER BY id DESC LIMIT 1"),{'u':user['id']}).mappings().one()
            self.assertEqual(row['status'],'Retry');self.assertEqual(row['attempts'],1)
        a.patch('/api/notification-preferences',json={'email_digest':False})
        self.assertFalse(a.get('/api/notification-preferences').json()['email_digest'])

    def test_oauth_state_and_explicit_linking(self):
        a,user=self.account()
        config={'GITHUB_CLIENT_ID':'test-client','GITHUB_CLIENT_SECRET':'test-secret'}
        with patch.dict(os.environ,config):
            self.assertEqual(a.get('/api/auth/oauth/providers').json()['github'],True)
            start=a.get('/api/auth/oauth/github/start?link=true');self.assertEqual(start.status_code,307)
            self.assertIn('code_challenge=',start.headers['location'])
            state=a.cookies.get('careernova_oauth')
            bad=a.get('/api/auth/oauth/github/callback?state=wrong&code=abc');self.assertIn('invalid_state',bad.headers['location'])
            # Restart because the failed flow intentionally clears its browser binding.
            a.get('/api/auth/oauth/github/start?link=true');state=a.cookies.get('careernova_oauth')
            class TokenResponse:
                def raise_for_status(self):pass
                def json(self):return {'access_token':'test-only'}
            with patch.object(oauth_routes.requests,'post',return_value=TokenResponse()),patch.object(oauth_routes,'provider_identity',return_value=('test-subject-'+str(user['id']),user['email'],'Tester')):
                linked=a.get('/api/auth/oauth/github/callback',params={'state':state,'code':'test-code'})
                self.assertEqual(linked.status_code,307);self.assertTrue(linked.headers['location'].endswith('/dashboard'))
                self.assertEqual(a.get('/api/auth/me').json()['id'],user['id'])
                self.assertEqual(a.get('/api/auth/oauth/identities').json()['providers'],['github'])
                a.cookies.set('careernova_oauth',state)
                self.assertIn('expired_sign_in',a.get('/api/auth/oauth/github/callback',params={'state':state,'code':'test-code'}).headers['location'])
        with backend.engine.connect() as conn:
            self.assertEqual(conn.execute(text('SELECT COUNT(*) FROM public.oauth_identities WHERE user_id=:u'),{'u':user['id']}).scalar_one(),1)

if __name__=='__main__':unittest.main(verbosity=2)
