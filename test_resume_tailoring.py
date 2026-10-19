import io
import secrets
import unittest
from docx import Document
from resume_tailoring import prepare_draft, make_docx


class TailoringTests(unittest.TestCase):
    def test_supported_skills_and_facts_only(self):
        source = 'Asha Candidate\nExperience\nAcme | 2022–2025\n- Built React interfaces.\n- Developed Python SQL tools for 100 users.\nEducation\nUniversity degree'
        result = prepare_draft(source, 'Build Python and SQL services using Docker and Kubernetes.')
        self.assertIn('Python', result['matched_skills'])
        self.assertIn('Docker', result['missing_skills'])
        self.assertNotIn('Docker',result['text'])
        for line in source.splitlines():
            self.assertIn(line,result['text'])
        self.assertLess(result['text'].index('- Developed'),result['text'].index('- Built'))

    def test_docx_round_trip(self):
        value = '\nAsha Candidate\nSkills\nPython, SQL\nExperience\nBuilt tools for 100 users.'
        doc = Document(io.BytesIO(make_docx(value)))
        self.assertEqual('\n'.join(p.text for p in doc.paragraphs), value)


class TailoringIntegration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import backend
        from fastapi.testclient import TestClient
        cls.backend=backend; cls.context=TestClient(backend.app); cls.context.__enter__(); cls.users=[]

    @classmethod
    def tearDownClass(cls):
        from sqlalchemy import text
        with cls.backend.engine.begin() as conn:
            for uid in cls.users:
                conn.execute(text('DELETE FROM public.users WHERE id=:u'),{'u':uid})
        cls.context.__exit__(None,None,None)

    def account(self):
        from fastapi.testclient import TestClient
        client=TestClient(self.backend.app,headers={'X-CareerNova-Client':'web'})
        response=client.post('/api/auth/register',json={'email':f'cn-tailor-{secrets.token_hex(8)}@example.com','full_name':'Resume Tester','password':'Tailor-Test-Only-2026!'})
        self.assertEqual(response.status_code,201,response.text)
        self.users.append(client.get('/api/auth/me').json()['id'])
        return client

    def test_private_draft_save_download_and_source_unchanged(self):
        a=self.account(); b=self.account()
        source='Asha Candidate\nExperience\n- Built Python SQL tools for 100 users.\nEducation\nUniversity degree'
        original=make_docx(source)
        upload=a.post('/api/resumes',data={'name':'Original'},files={'file':('original.docx',original)})
        self.assertEqual(upload.status_code,201,upload.text)
        rid=upload.json()['id']; payload={'resume_id':rid,'job_description':'Build Python and SQL services using Docker.'}
        self.assertEqual(b.post('/api/resume-tailor',json=payload).status_code,404)
        draft=a.post('/api/resume-tailor',json=payload)
        self.assertEqual(draft.status_code,200,draft.text)
        self.assertNotIn('Docker',draft.json()['text'])
        saved=a.post('/api/resume-tailor/save',json={'name':'Target role','text':draft.json()['text']})
        self.assertEqual(saved.status_code,201,saved.text)
        new_id=saved.json()['id']
        self.assertNotEqual(rid,new_id)
        self.assertEqual(a.get(f'/api/resumes/{rid}/download').content,original)
        downloaded=a.get(f'/api/resumes/{new_id}/download')
        doc=Document(io.BytesIO(downloaded.content))
        self.assertIn('Python','\n'.join(p.text for p in doc.paragraphs))
        self.assertEqual(b.get(f'/api/resumes/{new_id}/download').status_code,404)
        self.assertEqual(a.post('/api/resume-tailor',json={**payload,'job_description':'short'}).status_code,422)
