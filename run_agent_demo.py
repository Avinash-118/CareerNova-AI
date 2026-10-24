"""Extension-free local demo. Uses synthetic data and never submits a form."""
import io
import json
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from docx import Document
from application_session import ApplicationSession, validate_url

ROOT = Path(__file__).resolve().parent


class DemoHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path != '/test_application.html':
            self.send_error(404)
            return
        content = (ROOT / 'test_application.html').read_bytes()
        self.send_response(200)
        self.send_header('Content-Type', 'text/html; charset=utf-8')
        self.end_headers()
        self.wfile.write(content)

    def log_message(self, *_):
        pass


def main():
    server = ThreadingHTTPServer(('127.0.0.1', 8765), DemoHandler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    url = 'http://127.0.0.1:8765/test_application.html'
    validate_url(url)
    session = ApplicationSession(url)
    profile = {'first_name': 'Demo', 'last_name': 'Candidate',
               'email': 'demo@example.com', 'phone': '2025550147', 'location':'Boston'}
    document = Document()
    document.add_heading('Demo Candidate — Test Resume', 0)
    document.add_paragraph('Synthetic test data only. Skills: Python, SQL, React.')
    resume = io.BytesIO()
    document.save(resume)
    output = ROOT / '.deploy'
    output.mkdir(exist_ok=True)
    try:
        session.start()
        result = session.autofill(profile)
        # Select the resume upload question; leave experience and cover letter unanswered.
        index = next(i for i, field in enumerate(session.fields)
                     if field['type'] == 'file')
        session.revisit(index)
        session.upload('demo-resume.docx', resume.getvalue())
        def verify():
            page = session.page
            assert page.locator('#full_name').input_value() == 'Demo Candidate'
            assert page.locator('#email').input_value() == profile['email']
            assert page.locator('#phone').input_value() == profile['phone']
            assert page.locator('.select__single-value').inner_text() == 'Boston'
            assert page.locator('#resume').evaluate('(el) => el.files[0].name') == 'demo-resume.docx'
            page.screenshot(path=str(output / 'agent-demo-filled.png'), full_page=True)
        session.run(verify)
        report = {'profile_fields_filled': result['autofilled_count'],
                  'resume_uploaded': True, 'submitted': False,
                  'url': url, 'synthetic_data_only': True}
        (output / 'agent-demo-result.json').write_text(json.dumps(report, indent=2))
        print(json.dumps(report), flush=True)
        print('Demo browser stays open for up to 10 minutes. Close it when finished.', flush=True)
        for _ in range(600):
            if not session.run(lambda: session.browser.is_connected()):
                break
            time.sleep(1)
    finally:
        session.close()
        server.shutdown()
        server.server_close()


if __name__ == '__main__':
    main()
